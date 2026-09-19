from datetime import date, datetime
import hashlib
import json
import os
import re
from pathlib import Path
import sys
import tempfile
import urllib.parse
import urllib.request

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(os.environ.get("PROMOS_SOURCE_XLSX", ROOT.parent / "Credit Card Promos.xlsx"))
OUTPUT = Path(os.environ.get("PROMOS_DATA_OUTPUT", ROOT / "app" / "data"))
IMAGE_CACHE_BANKS = {"BPI", "RCBC"}
IMAGE_URL_OVERRIDES = {
    "https://www.bpi.com.ph/content/dam/bau/promos/2026-promos/hue,-laud,-wright,-the-aurora/content_card_hue_boracay.jpg_boracay.png":
        "https://www.bpi.com.ph/content/dam/bau/promos/2026-promos/hue,-laud,-wright,-the-aurora/content_card_hue_boracay.jpg",
}
CATEGORY_ORDER = [
    "Dining",
    "Shopping",
    "Travel & Leisure",
    "Installments & Financing",
    "Online",
    "Services",
    "Health & Wellness",
    "Welcome Gift",
    "Other",
]


def scalar(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return "" if value is None else str(value).strip()


def is_current_end_date(value):
    if not value:
        return True
    try:
        return date.fromisoformat(value[:10]) >= date.today()
    except ValueError:
        return True


def normalize_categories(raw_categories, promo, summary):
    raw_text = " ".join(raw_categories).lower()
    categories = []

    def add(category):
        if category not in categories:
            categories.append(category)

    if any(term in raw_text for term in ("dine", "dining", "restaurant")):
        add("Dining")
    if any(term in raw_text for term in ("shopping", "shop", "essentials", "technology")):
        add("Shopping")
    if any(term in raw_text for term in ("travel", "leisure", "stay", "entertainment")):
        add("Travel & Leisure")
    if any(term in raw_text for term in ("installment", "buy now pay later")):
        add("Installments & Financing")
    if any(term in raw_text for term in ("online", "e-commerce", "ecommerce")):
        add("Online")
    if "service" in raw_text:
        add("Services")
    if any(term in raw_text for term in ("health", "wellness", "beauty")):
        add("Health & Wellness")
    if "welcome" in raw_text:
        add("Welcome Gift")

    if not categories:
        text = promo.lower()
        if any(term in text for term in ("dine", "dining", "restaurant", "food", "cafe", "coffee", "ramen")):
            add("Dining")
        if any(term in text for term in ("shopping", "shop", "store", "mall", "grocery", "fashion", "technology", "appliance", "electronics", "jewelry", "miniso", "lazada", "shopee", "zara", "sperry", "keds", "merrell")):
            add("Shopping")
        if any(term in text for term in ("travel", "hotel", "resort", "stay", "flight", "airline", "agoda", "booking", "trip", "tour", "vacation", "holiday", "airport", "entertainment")):
            add("Travel & Leisure")
        if any(term in text for term in ("installment", "buy now pay later", "pay later", "flexipay", "cash2go", "0%")):
            add("Installments & Financing")
        if any(term in text for term in ("online", "e-commerce", "ecommerce", "foodpanda", "qr ph", "digital", "bdo pay", "payeasy")):
            add("Online")
        if any(term in text for term in ("service", "insurance", "automotive", "car care", "fuel", "gas station", "utility")):
            add("Services")
        if any(term in text for term in ("health", "wellness", "beauty", "diagnostic", "clinic", "dental", "skin", "spa", "salon", "pharmacy", "optical", "fitness")):
            add("Health & Wellness")
        if any(term in text for term in ("welcome", "naffl", "new card", "member-get-member", "referral", "acquisition")):
            add("Welcome Gift")
    if not categories:
        add("Other")

    return sorted(categories, key=CATEGORY_ORDER.index)


def cache_bank_image(bank, url):
    if not url.startswith(("http://", "https://")):
        return ""

    parsed = urllib.parse.urlsplit(url)
    extension = Path(urllib.parse.unquote(parsed.path)).suffix.lower()
    if extension not in {".jpg", ".jpeg", ".png", ".webp", ".gif"}:
        extension = ".jpg"
    filename = f"{hashlib.sha1(url.encode()).hexdigest()[:16]}{extension}"
    cache_dir = ROOT / "public" / "promo-images" / bank.lower()
    target = cache_dir / filename

    if not target.exists() or target.stat().st_size == 0:
        cache_dir.mkdir(parents=True, exist_ok=True)
        encoded_path = urllib.parse.quote(
            urllib.parse.unquote(parsed.path),
            safe="/:@-._~!$&'()*+,;=",
        )
        request_url = urllib.parse.urlunsplit(
            (parsed.scheme, parsed.netloc, encoded_path, parsed.query, parsed.fragment)
        )
        try:
            request = urllib.request.Request(request_url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(request, timeout=30) as response:
                payload = response.read()
            if not payload:
                raise RuntimeError("empty response")
            with tempfile.NamedTemporaryFile(dir=cache_dir, delete=False) as temporary:
                temporary.write(payload)
                temporary_path = Path(temporary.name)
            temporary_path.replace(target)
        except Exception as error:
            print(f"Warning: could not cache {bank} image {url}: {error}", file=sys.stderr)
            return url

    return f"/promo-images/{bank.lower()}/{filename}"


def write_json(path, payload, **kwargs):
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, **kwargs), encoding="utf-8")
    temporary.replace(path)


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(SOURCE, data_only=True)
    sheet = workbook["All Promos"]
    headers = [scalar(sheet.cell(row=6, column=column).value) for column in range(1, sheet.max_column + 1)]
    columns = {header: index + 1 for index, header in enumerate(headers)}

    def value(row, header):
        cell = sheet.cell(row=row, column=columns[header])
        if cell.hyperlink and cell.hyperlink.target:
            return cell.hyperlink.target
        return scalar(cell.value)

    promos = []
    for row in range(7, sheet.max_row + 1):
        if not any(sheet.cell(row=row, column=column).value is not None for column in range(1, sheet.max_column + 1)):
            continue

        bank = value(row, "Bank")
        promo = value(row, "Promo")
        offer_url = value(row, "Offer page")
        end_date = value(row, "End date") or None
        if end_date and not is_current_end_date(end_date):
            continue
        categories_raw = value(row, "Category")
        raw_categories = [part.strip() for part in re.split(r"\s*;\s*", categories_raw) if part.strip()]
        summary = value(row, "Offer summary")
        categories = normalize_categories(raw_categories, promo, summary)
        identifier = hashlib.sha1(f"{bank}|{promo}|{offer_url}|{row}".encode()).hexdigest()[:16]

        image_url = value(row, "Image")
        image_url = image_url if image_url.startswith(("http://", "https://")) else ""
        image_url = IMAGE_URL_OVERRIDES.get(image_url, image_url)
        if bank in IMAGE_CACHE_BANKS and image_url:
            image_url = cache_bank_image(bank, image_url)

        promos.append({
            "id": identifier,
            "bank": bank,
            "promo": promo,
            "category": "; ".join(categories),
            "categories": categories,
            "summary": summary,
            "startDate": value(row, "Start date") or None,
            "endDate": end_date,
            "dateCheck": value(row, "Date check"),
            "cardTypes": value(row, "Card types (listing)"),
            "offerUrl": offer_url if offer_url.startswith(("http://", "https://")) else "",
            "imageUrl": image_url,
            "originalDateWording": value(row, "Original date wording"),
            "checkedDate": value(row, "Checked date"),
            "dateAdded": value(row, "Date added") or None,
            "sourceRow": row,
        })

    metadata = {
        "sourceFile": SOURCE.name,
        "sheet": "All Promos",
        "checkedAt": scalar(sheet.cell(row=3, column=2).value),
        "listedOffers": len(promos),
        "banks": sorted({promo["bank"] for promo in promos if promo["bank"]}),
        "categories": sorted({category for promo in promos for category in promo["categories"]}),
    }

    client_promos = [
        {
            "id": promo["id"],
            "bank": promo["bank"],
            "promo": promo["promo"],
            "category": promo["category"],
            "categories": promo["categories"],
            "summary": promo["summary"],
            "startDate": promo["startDate"],
            "endDate": promo["endDate"],
            "dateCheck": promo["dateCheck"],
            "offerUrl": promo["offerUrl"],
            "imageUrl": promo["imageUrl"],
            "originalDateWording": promo["originalDateWording"],
            "dateAdded": promo["dateAdded"],
        }
        for promo in promos
    ]

    write_json(OUTPUT / "promos.json", promos, ensure_ascii=True, separators=(",", ":"))
    write_json(OUTPUT / "promos-client.json", client_promos, ensure_ascii=True, separators=(",", ":"))
    write_json(OUTPUT / "meta.json", metadata, ensure_ascii=True, indent=2)
    print(f"Exported {len(promos)} promotions from {SOURCE.name}")


if __name__ == "__main__":
    main()
